"""SOP → Markdown → sections (docs/planning/spec-sop-section-grounding.md). Synthetic text only:
this repo is public, the real client SOPs never appear here."""

import io

import httpx
import pytest

from app.services import sop_markdown
from app.services.sop_markdown import _analyze as real_analyze
from app.sop.docx_markdown import docx_to_markdown
from app.sop.sections import full_text, parse_sections

# Shaped like Document Intelligence's Markdown of a numbered SOP: running page headers, page
# breaks, numbered headings, and clauses whose number sits alone on its line.
DI_MARKDOWN = """<!-- PageHeader="Number: DOC-001" -->

# STANDARD OPERATING PROCEDURE

<!-- PageNumber="Page 1 of 3" -->

Company Confidential

## 1. PURPOSE

This procedure describes how the widget is inspected.

## 2. RESPONSIBILITIES

2.1
Inspector:

2.1.1
Checks every widget before release.

<!-- PageBreak -->

# STANDARD OPERATING PROCEDURE

Company Confidential

2.1.2 Records the result in the log.

2.2 Supervisor: reviews the log weekly.

<table>
<tr><td>Role</td><td>Task</td></tr>
</table>

<!-- PageBreak -->

# STANDARD OPERATING PROCEDURE

Company Confidential

## 3. RECORDS

Logs are kept for 3.5 years.

Company Confidential
"""


def _by_number(sections):
    return {s.number: i for i, s in enumerate(sections)}


def test_numbered_clauses_nest_and_running_headers_are_dropped_after_their_first_copy():
    sections = parse_sections(DI_MARKDOWN)
    numbers = [s.number for s in sections]
    assert numbers == ["§0", "1", "2", "2.1", "2.1.1", "2.1.2", "2.2", "3"]
    idx = _by_number(sections)
    assert sections[idx["2.1"]].title == "Inspector:"
    assert sections[idx["2.1.1"]].parent == idx["2.1"]
    assert sections[idx["2.1.2"]].page_start == 2
    # The running title and footer survive once (in the preamble), never inside a clause.
    assert "STANDARD OPERATING PROCEDURE" in sections[0].text
    assert all("Company Confidential" not in s.text for s in sections[1:])
    # "3.5 years" is prose, not a clause 3.5.
    assert "3.5" not in numbers


def test_a_full_section_is_the_whole_passage_with_every_subsection_and_table():
    sections = parse_sections(DI_MARKDOWN)
    text = full_text(sections, _by_number(sections)["2"])
    for part in (
        "2 RESPONSIBILITIES",
        "2.1 Inspector:",
        "Checks every widget",
        "2.1.2 Records",
        "2.2 Supervisor",
        "<table>",
        "Role",
    ):
        assert part in text
    assert "PURPOSE" not in text and "RECORDS" not in text


def test_unnumbered_headings_split_and_a_preamble_is_kept():
    md = "Intro before any heading.\n\n# Overview\n\nText A.\n\n## Detail\n\nText B.\n"
    sections = parse_sections(md)
    assert [(s.number, s.title, s.level) for s in sections] == [
        ("§0", "", 1),
        ("§1", "Overview", 1),
        ("§2", "Detail", 2),
    ]
    assert sections[0].text == "Intro before any heading."
    assert full_text(sections, 1) == "Overview\nText A.\n\nDetail\nText B."


def test_a_document_with_no_heading_is_one_section():
    (only,) = parse_sections("Just a paragraph.\n\nAnd another.")
    assert (only.number, only.text) == ("§0", "Just a paragraph.\nAnd another.")


def test_a_skipped_clause_does_not_orphan_the_ones_after_it():
    md = "1 Scope\n\n1.1 First.\n\n1.3 Third (1.2 was not split out).\n\n1.4 Fourth.\n\n2 Next\n"
    assert [s.number for s in parse_sections(md)] == ["1", "1.1", "1.3", "1.4", "2"]


def _docx(build) -> bytes:
    import docx

    document = docx.Document()
    build(document)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def test_word_headings_come_from_styles_caps_and_bold_and_tables_become_markdown():
    def build(d):
        d.add_paragraph("Preface text.")
        d.add_heading("Styled heading", level=1)
        d.add_paragraph("Under the styled heading.")
        d.add_paragraph("TYPED IN CAPITALS")
        d.add_paragraph("Under the caps heading.")
        bold = d.add_paragraph()
        bold.add_run("Bold subheading").bold = True
        d.add_paragraph("Under the bold heading.")
        table = d.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Visit", "Window"
        table.cell(1, 0).text, table.cell(1, 1).text = "SIV", "Day 1"

    md = docx_to_markdown(_docx(build))
    assert "# Styled heading" in md and "# TYPED IN CAPITALS" in md and "## Bold subheading" in md
    assert "| Visit | Window |" in md and "| SIV | Day 1 |" in md
    sections = parse_sections(md)
    assert [s.title for s in sections] == [
        "",
        "Styled heading",
        "TYPED IN CAPITALS",
        "Bold subheading",
    ]
    assert "| SIV | Day 1 |" in full_text(sections, 3)


# --- conversion is all or nothing --------------------------------------------------------------


def _pages(*texts):
    return list(texts)


def test_page_gaps_flags_a_page_whose_recurring_words_are_missing():
    page1 = (
        "Widget inspection record checklist release batch quality signed dated VV-QDOC-12345 " * 2
    )
    page2 = (
        "Supervisor review logbook entries weekly monthly escalation deviation closure audit " * 2
    )
    assert sop_markdown.page_gaps(_pages(page1, page2), page1 + page2) == []
    gaps = sop_markdown.page_gaps(_pages(page1, page2), page1)
    assert [g[0] for g in gaps] == [2]
    # Rotated text comes out of the text layer reversed; merged words lose their spaces.
    assert (
        sop_markdown.page_gaps(
            _pages("reganaM " * 12 + "ataD tnemeganaM " * 6), "Manager Data Management"
        )
        == []
    )


@pytest.mark.asyncio
async def test_a_pdf_whose_pages_dont_all_come_back_fails_with_no_markdown(monkeypatch):
    monkeypatch.setattr(sop_markdown.get_settings(), "azure_foundry_endpoint", "https://di.example")
    monkeypatch.setattr(sop_markdown, "_pdf_page_texts", lambda _c: ["one " * 20, "two " * 20])

    async def two_of_three(_content, _endpoint, page_count):
        assert page_count == 2  # every page is named explicitly
        return {"content": "one", "pages": [{}]}

    monkeypatch.setattr(sop_markdown, "_analyze", two_of_three)
    result = await sop_markdown.to_markdown(b"%PDF", "x.pdf")
    assert (result.markdown, result.source) == ("", "failed")
    assert "1 of 2 pages converted" in result.error


@pytest.mark.asyncio
async def test_a_pdf_missing_a_pages_content_fails_and_a_complete_one_succeeds(monkeypatch):
    texts = [
        "Alpha bravo charlie delta echo foxtrot golf hotel india juliet " * 2,
        "Kilo lima mike november oscar papa quebec romeo sierra tango " * 2,
    ]
    monkeypatch.setattr(sop_markdown.get_settings(), "azure_foundry_endpoint", "https://di.example")
    monkeypatch.setattr(sop_markdown, "_pdf_page_texts", lambda _c: texts)

    async def partial(*_a):
        return {"content": texts[0], "pages": [{}, {}]}

    monkeypatch.setattr(sop_markdown, "_analyze", partial)
    failed = await sop_markdown.to_markdown(b"%PDF", "x.pdf")
    assert failed.source == "failed" and "page 2" in failed.error

    async def whole(*_a):
        return {"content": "\n".join(texts), "pages": [{}, {}]}

    monkeypatch.setattr(sop_markdown, "_analyze", whole)
    ok = await sop_markdown.to_markdown(b"%PDF", "x.pdf")
    assert (ok.source, ok.error) == ("document_intelligence", "")


@pytest.mark.asyncio
async def test_without_document_intelligence_the_text_layer_is_used_and_labelled(monkeypatch):
    monkeypatch.setattr(sop_markdown.get_settings(), "azure_foundry_endpoint", "")
    monkeypatch.setattr(sop_markdown, "_pdf_page_texts", lambda _c: ["p1", "p2"])
    result = await sop_markdown.to_markdown(b"%PDF", "x.pdf")
    assert (result.source, result.markdown) == ("pdf_text", "p1\n\n<!-- PageBreak -->\n\np2")
    unsupported = await sop_markdown.to_markdown(b"x", "x.xyz")
    assert unsupported.source == "failed" and "unsupported" in unsupported.error


def test_a_page_is_not_covered_by_its_words_appearing_on_a_distant_page():
    words = "Widget inspection record checklist release batch quality signed dated archive "
    pages = _pages(words * 2, "filler " * 30, "padding " * 30, words * 2)
    br = "\n<!-- PageBreak -->\n"
    whole = br.join([words, "filler", "padding", words])
    assert sop_markdown.page_gaps(pages, whole) == []
    # Page 4 lost its content; the same words on page 1 (three pages away) do not cover it.
    lost = br.join([words, "filler", "padding", ""])
    assert [g[0] for g in sop_markdown.page_gaps(pages, lost)] == [4]


def test_chinese_pages_are_checked_by_character_pairs():
    page = "监查访视报告必须在访视后五个工作日内完成并提交审核" * 2
    assert sop_markdown.page_gaps(_pages(page), page) == []
    assert [g[0] for g in sop_markdown.page_gaps(_pages(page), "监查访视")] == [1]


def test_word_auto_numbering_becomes_the_clause_number():
    import docx

    d = docx.Document()
    for title in ("PURPOSE", "SCOPE", "RESPONSIBILITIES"):
        d.add_paragraph(title, style="List Number")
        d.add_paragraph(f"Body text for {title.lower()}, long enough to read as a paragraph.")
    buf = io.BytesIO()
    d.save(buf)
    sections = parse_sections(docx_to_markdown(buf.getvalue()))
    assert [(s.number, s.title) for s in sections] == [
        ("1", "PURPOSE"),
        ("2", "SCOPE"),
        ("3", "RESPONSIBILITIES"),
    ]


@pytest.mark.asyncio
async def test_a_conversion_with_no_text_is_a_failure(monkeypatch):
    monkeypatch.setattr(sop_markdown.get_settings(), "azure_foundry_endpoint", "")
    monkeypatch.setattr(sop_markdown, "_pdf_page_texts", lambda _c: ["", ""])
    scanned = await sop_markdown.to_markdown(b"%PDF", "scan.pdf")
    assert (scanned.source, scanned.error) == ("failed", "the conversion produced no text")
    empty = await sop_markdown.to_markdown(b"  \n", "empty.md")
    assert empty.source == "failed"


@pytest.mark.asyncio
async def test_an_unexpected_converter_error_is_a_failed_conversion(monkeypatch):
    monkeypatch.setattr(sop_markdown.get_settings(), "azure_foundry_endpoint", "https://di.example")

    def broken(_content):
        raise ValueError("not a PDF")

    monkeypatch.setattr(sop_markdown, "_pdf_page_texts", broken)
    result = await sop_markdown.to_markdown(b"junk", "x.pdf")
    assert (result.source, result.error) == ("failed", "ValueError: not a PDF")


def _di_client(monkeypatch, handler):
    real_client = httpx.AsyncClient

    def client(**kwargs):
        return real_client(transport=httpx.MockTransport(handler), **kwargs)

    async def token(_scope):
        return "tok"

    monkeypatch.setattr(sop_markdown.httpx, "AsyncClient", client)
    monkeypatch.setattr(sop_markdown, "get_bearer_token", token)
    monkeypatch.setattr(sop_markdown, "DI_POLL_SECONDS", 0)


@pytest.mark.asyncio
async def test_document_intelligence_is_asked_for_every_page_and_throttling_is_waited_out(
    monkeypatch,
):
    calls: list[str] = []
    polls = iter(["running", "succeeded"])

    def handler(request: httpx.Request) -> httpx.Response:
        calls.append(f"{request.method} {request.url}")
        if request.method == "POST":
            if len(calls) == 1:
                return httpx.Response(429, headers={"retry-after": "0"})
            return httpx.Response(202, headers={"operation-location": "https://di.example/op/1"})
        body = {"status": next(polls), "analyzeResult": {"content": "md", "pages": [{}, {}, {}]}}
        return httpx.Response(200, json=body)

    _di_client(monkeypatch, handler)
    result = await real_analyze(b"%PDF", "https://di.example/", 3)
    assert result["content"] == "md"
    assert "pages=1-3" in calls[0] and "features=ocrHighResolution" in calls[0]
    assert [c.split()[0] for c in calls] == ["POST", "POST", "GET", "GET"]


@pytest.mark.asyncio
async def test_a_failed_analysis_and_a_missing_token_raise(monkeypatch):
    def handler(request: httpx.Request) -> httpx.Response:
        if request.method == "POST":
            return httpx.Response(202, headers={"operation-location": "https://di.example/op/1"})
        return httpx.Response(200, json={"status": "failed", "error": {"code": "BadPdf"}})

    _di_client(monkeypatch, handler)
    with pytest.raises(RuntimeError, match="BadPdf"):
        await real_analyze(b"%PDF", "https://di.example", 1)

    async def no_token(_scope):
        return ""

    monkeypatch.setattr(sop_markdown, "get_bearer_token", no_token)
    with pytest.raises(RuntimeError, match="No Entra token"):
        await real_analyze(b"%PDF", "https://di.example", 1)


def _docx_with_table(rows: list[list[str]], merge_first_row: bool = False) -> bytes:
    import docx

    d = docx.Document()
    table = d.add_table(rows=len(rows), cols=len(rows[0]))
    for r, values in enumerate(rows):
        for c, value in enumerate(values):
            table.cell(r, c).text = value
    if merge_first_row:
        table.cell(0, 0).merge(table.cell(0, len(rows[0]) - 1)).text = rows[0][0]
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_a_merged_cell_is_written_once_and_columns_still_line_up():
    md = docx_to_markdown(
        _docx_with_table(
            [["Study visits", "", ""], ["Visit", "Window", "Owner"], ["V1", "Day 1", "CRA"]],
            merge_first_row=True,
        )
    )
    assert md.count("Study visits") == 1
    assert md.splitlines()[0] == "| Study visits |  |  |"
    assert "| V1 | Day 1 | CRA |" in md


def test_a_form_table_becomes_labelled_sections():
    import docx

    d = docx.Document()
    table = d.add_table(rows=3, cols=2)
    rows = [
        "Job Description",
        "General Description: Leads the study team.",
        "Essential Functions: Plans visits.",
    ]
    for r, text in enumerate(rows):
        table.cell(r, 0).merge(table.cell(r, 1)).text = text
    buf = io.BytesIO()
    d.save(buf)
    sections = parse_sections(docx_to_markdown(buf.getvalue()))
    assert [(s.title, s.text) for s in sections] == [
        ("", "Job Description"),
        ("General Description", "Leads the study team."),
        ("Essential Functions", "Plans visits."),
    ]


def test_a_one_column_table_and_a_time_are_not_a_form():
    md = docx_to_markdown(_docx_with_table([["10:30 site review"], ["Note: bring the log"]]))
    assert md.splitlines()[0] == "| 10:30 site review |"
    assert "## " not in md


async def test_a_powerpoint_is_converted_by_document_intelligence_slide_by_slide(monkeypatch):
    """PowerPoint goes to DI (owner, 2026-10-09): every slide title a section, the slide its page.
    No page count is asked for and no high-resolution OCR (a deck is not a scan)."""
    monkeypatch.setattr(sop_markdown.get_settings(), "azure_foundry_endpoint", "https://di.example")
    seen = {}

    async def deck(content, endpoint, page_count, content_type=sop_markdown._PDF):
        seen.update(page_count=page_count, content_type=content_type)
        return {
            "content": "# 1 PURPOSE\n\nCovers widgets.\n\n<!-- PageBreak -->\n\n"
            "# 2 RELEASE\n\nSigned within 24 hours.\n",
            "pages": [{}, {}],
        }

    monkeypatch.setattr(sop_markdown, "_analyze", deck)
    result = await sop_markdown.to_markdown(b"PK", "Release.pptx")
    assert result.source == "document_intelligence"
    assert seen == {"page_count": None, "content_type": sop_markdown._OFFICE[".pptx"]}
    sections = parse_sections(result.markdown)
    assert [(s.number, s.page_start) for s in sections] == [("1", 1), ("2", 2)]


async def test_a_powerpoint_without_document_intelligence_is_a_labelled_failure(monkeypatch):
    monkeypatch.setattr(sop_markdown.get_settings(), "azure_foundry_endpoint", "")
    result = await sop_markdown.to_markdown(b"PK", "Release.pptx")
    assert result.source == "failed" and "Document Intelligence" in result.error


async def test_excel_is_not_a_supported_sop_format(monkeypatch):
    monkeypatch.setattr(sop_markdown.get_settings(), "azure_foundry_endpoint", "https://di.example")
    result = await sop_markdown.to_markdown(b"PK", "Plan.xlsx")
    assert result.source == "failed" and "unsupported" in result.error


def test_a_heading_deeper_than_six_hashes_is_still_a_section():
    """DI writes a seventh-level Word heading as "#######" (a client SOP's "14. SIGNATURES" was
    lost this way, 2026-10-09). It is a heading, kept at level 6."""
    sections = parse_sections("## 1. RECORDS\n\nKept.\n\n####### 2. SIGNATURES\n\nSigned.\n")
    assert [s.number for s in sections] == ["1", "2"]
    assert sections[1].title == "SIGNATURES"


def _word(*paragraphs: str) -> bytes:
    import docx

    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


WORD_TEXT = [
    "Inspector checks every widget before release and records the widget result.",
    "Supervisor reviews the log weekly; the supervisor signs it weekly.",
]


@pytest.mark.asyncio
async def test_word_goes_through_libreoffice_to_document_intelligence(monkeypatch):
    """Word → PDF (LibreOffice) → DI (owner, 2026-10-09). Completeness is checked against the
    .docx's own text: a DI reading that drops a paragraph fails the whole document."""
    monkeypatch.setattr(sop_markdown.get_settings(), "azure_foundry_endpoint", "https://di.example")
    monkeypatch.setattr(sop_markdown, "libreoffice", lambda: "/usr/bin/soffice")
    monkeypatch.setattr(sop_markdown, "_docx_to_pdf", lambda content, soffice: b"%PDF-word")
    monkeypatch.setattr(sop_markdown, "_pdf_page_texts", lambda _c: ["page one", "page two"])
    seen = {}

    async def read(content, endpoint, page_count, content_type=sop_markdown._PDF):
        seen.update(content=content, page_count=page_count, content_type=content_type)
        return {"content": "# 1 RELEASE\n\n" + "\n\n".join(WORD_TEXT), "pages": [{}, {}]}

    monkeypatch.setattr(sop_markdown, "_analyze", read)
    ok = await sop_markdown.to_markdown(_word(*WORD_TEXT), "Release.docx")
    assert (ok.source, ok.error) == ("document_intelligence", "")
    # The PDF went to DI as a PDF, every page named.
    assert seen == {"content": b"%PDF-word", "page_count": 2, "content_type": sop_markdown._PDF}

    async def half(*_a, **_k):
        return {"content": WORD_TEXT[0], "pages": [{}, {}]}

    monkeypatch.setattr(sop_markdown, "_analyze", half)
    failed = await sop_markdown.to_markdown(_word(*WORD_TEXT), "Release.docx")
    assert failed.source == "failed" and "not fully read" in failed.error


@pytest.mark.asyncio
async def test_without_libreoffice_word_falls_back_only_where_it_is_not_required(monkeypatch):
    settings = sop_markdown.get_settings()
    monkeypatch.setattr(settings, "azure_foundry_endpoint", "https://di.example")
    monkeypatch.setattr(sop_markdown, "libreoffice", lambda: None)
    # Dev / CI: our own converter, labelled.
    monkeypatch.setattr(settings, "sop_require_libreoffice", False)
    dev = await sop_markdown.to_markdown(_word(*WORD_TEXT), "Release.docx")
    assert dev.source == "docx" and "Inspector checks" in dev.markdown
    # A published image requires LibreOffice: no silent fallback, a failure that says why.
    monkeypatch.setattr(settings, "sop_require_libreoffice", True)
    image = await sop_markdown.to_markdown(_word(*WORD_TEXT), "Release.docx")
    assert image.source == "failed" and "LibreOffice is not installed" in image.error
    # ... nor without Document Intelligence.
    monkeypatch.setattr(sop_markdown, "libreoffice", lambda: "/usr/bin/soffice")
    monkeypatch.setattr(settings, "azure_foundry_endpoint", "")
    no_di = await sop_markdown.to_markdown(_word(*WORD_TEXT), "Release.docx")
    assert no_di.source == "failed" and "Document Intelligence" in no_di.error


@pytest.mark.asyncio
async def test_an_empty_pdf_from_libreoffice_is_a_clear_failure(monkeypatch):
    monkeypatch.setattr(sop_markdown.get_settings(), "azure_foundry_endpoint", "https://di.example")
    monkeypatch.setattr(sop_markdown, "libreoffice", lambda: "/usr/bin/soffice")
    monkeypatch.setattr(sop_markdown, "_docx_to_pdf", lambda content, soffice: b"%PDF")
    monkeypatch.setattr(sop_markdown, "_pdf_page_texts", lambda _c: [])
    result = await sop_markdown.to_markdown(_word(*WORD_TEXT), "Release.docx")
    assert result.source == "failed" and "empty PDF" in result.error


def test_a_libreoffice_that_hangs_is_killed_with_its_children(monkeypatch, tmp_path):
    """soffice starts soffice.bin; a timeout kills the whole process group."""
    hang = tmp_path / "soffice"
    hang.write_text("#!/bin/sh\nsleep 30 &\nwait\n")
    hang.chmod(0o755)
    monkeypatch.setattr(sop_markdown, "LIBREOFFICE_TIMEOUT_SECONDS", 0.5)
    with pytest.raises(RuntimeError, match="took over"):
        sop_markdown._docx_to_pdf(b"PK", str(hang))


def test_text_gaps_names_the_recurring_words_the_markdown_lost():
    coverage, missing = sop_markdown.text_gaps("\n".join(WORD_TEXT), WORD_TEXT[0])
    assert coverage < sop_markdown.PAGE_COVERAGE_MIN and "supervisor" in missing
    assert sop_markdown.text_gaps("\n".join(WORD_TEXT), "\n".join(WORD_TEXT)) == (1.0, [])


@pytest.mark.skipif(sop_markdown.libreoffice() is None, reason="LibreOffice is not installed")
def test_libreoffice_prints_a_word_file_to_pdf():
    pdf = sop_markdown._docx_to_pdf(_word(*WORD_TEXT), sop_markdown.libreoffice())
    assert pdf.startswith(b"%PDF")
    assert "Inspector checks" in "".join(sop_markdown._pdf_page_texts(pdf))
