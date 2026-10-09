"""Fail the image build unless LibreOffice prints a Word document to PDF (run by the Dockerfile).

Uses the same code path as an SOP upload (``sop_markdown._docx_to_pdf``), on a .docx with a
Chinese and an English paragraph: the PDF must carry the text, the Chinese drawn in a CJK font.
"""

import io
import subprocess
import sys

import docx
import pdfplumber

from app.services.sop_markdown import _docx_to_pdf, libreoffice


def main() -> int:
    soffice = libreoffice()
    if not soffice:
        print("check_libreoffice: soffice not found", file=sys.stderr)
        return 1
    # A Chinese font must be installed, or LibreOffice lays Chinese out with a substitute.
    chinese = subprocess.run(["fc-list", ":lang=zh"], capture_output=True, text=True).stdout
    if not chinese.strip():
        print("check_libreoffice: no Chinese font installed", file=sys.stderr)
        return 1
    document = docx.Document()
    document.add_heading("1 Purpose", level=1)
    document.add_paragraph("Report every serious adverse event within 24 hours.")
    document.add_paragraph("严重不良事件须在24小时内报告。")
    buffer = io.BytesIO()
    document.save(buffer)
    pdf = _docx_to_pdf(buffer.getvalue(), soffice)
    # The Chinese must be DRAWN in a Chinese font (Document Intelligence reads the page, not the
    # PDF's text layer): without one, LibreOffice falls back to a Latin font with no Chinese
    # glyphs (seen on macOS: "LinuxLibertineG") and the page shows empty boxes.
    with pdfplumber.open(io.BytesIO(pdf)) as reader:
        chars = [c for page in reader.pages for c in page.chars]
    latin = "".join(c["text"] for c in chars if c["text"].isascii())
    fonts = {c["fontname"] for c in chars if not c["text"].isascii()}
    if "adverse event" not in latin or not fonts or not all("CJK" in f for f in fonts):
        print(f"check_libreoffice: Chinese drawn in {fonts}: {latin[:120]!r}", file=sys.stderr)
        return 1
    print(f"check_libreoffice: OK ({soffice}, {len(pdf)} bytes)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
